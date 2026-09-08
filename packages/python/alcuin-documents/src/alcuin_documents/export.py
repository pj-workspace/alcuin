"""Editable DOCX export of text Artifacts, without loading external resources."""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from io import BytesIO
import json
import re
from typing import Any
from urllib.parse import quote, urlsplit

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph
from markdown_it import MarkdownIt
from markdown_it.rules_inline.state_inline import StateInline
from markdown_it.token import Token


DOCX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
MAX_EXPORT_CHARS = 1_000_000
_CITATION = re.compile(r"\[\[cite:([A-Za-z0-9][A-Za-z0-9._:-]{0,199})\]\]")
_CITATION_ID = re.compile(r"[A-Za-z0-9][A-Za-z0-9._:-]{0,199}")
_INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
_PORTABLE_CITATION = re.compile(
    r"\[\[cite:([A-Za-z0-9][A-Za-z0-9._:-]{0,199})\]\]"
    r"|\[[^\]\n]*\]\(alcuin-citation:([A-Za-z0-9][A-Za-z0-9._:-]{0,199})\)"
)


class DocumentExportError(ValueError):
    """The Artifact cannot be exported through the supported text document path."""


def export_artifact_markdown(
    content: str, *, citations: Sequence[Mapping[str, Any]] = ()
) -> str:
    """Keep Markdown source intact while making exact evidence markers portable."""
    if len(content) > MAX_EXPORT_CHARS:
        raise DocumentExportError("Artifact exceeds the document export limit")
    registry = _Citations(citations)

    def escape(value: object) -> str:
        single_line = re.sub(r"\s+", " ", _clean(value)).strip()
        return re.sub(r"([\\`*_{}\[\]<>])", r"\\\1", single_line)

    def marker(match: re.Match[str]) -> str:
        identifier = match.group(1) or match.group(2)
        record = registry.records.get(identifier)
        if record is None:
            return registry.unresolved(identifier)
        registry.used.add(identifier)
        number = registry.numbers[identifier]
        href = _safe_link(str(record.get("locator") or record.get("url") or ""))
        if href:
            return f"[{number}](<{quote(href, safe='/:?&=#@!$%()*+,;[]')}>)"
        return f"[{number}]"

    def inline(value: str) -> str:
        # Let CommonMark identify actual text. A global substitution would corrupt
        # code, existing link destinations/labels, escaped syntax or HTML attributes.
        replacements: list[tuple[int, int, str]] = []
        state = StateInline(value, parser, environment, [])

        def capture(current: StateInline, silent: bool) -> bool:
            # Link-label lookahead must keep CommonMark's bracket nesting intact.
            if silent:
                return False
            match = _PORTABLE_CITATION.match(current.src, current.pos, current.posMax)
            if match is None:
                return False
            if current is state and current.linkLevel == 0:
                replacements.append((match.start(), match.end(), marker(match)))
            token = current.push("text", "", 0)
            token.content = match.group()
            current.pos = match.end()
            return True

        parser.inline.ruler.at("alcuin_citation", capture)
        parser.inline.tokenize(state)
        parts: list[str] = []
        cursor = 0
        for start, end, replacement in replacements:
            parts.extend((value[cursor:start], replacement))
            cursor = end
        parts.append(value[cursor:])
        return "".join(parts)

    lines = content.splitlines(keepends=True)
    parser = MarkdownIt("commonmark")
    environment: dict[str, Any] = {}
    blocks = parser.parse(content, environment)
    parser.inline.ruler.before("link", "alcuin_citation", lambda _state, _silent: False)
    parts: list[str] = []
    cursor = 0
    for token in blocks:
        if token.type != "inline" or not token.map:
            continue
        start, end = token.map
        if start < cursor:
            continue
        parts.append("".join(lines[cursor:start]))
        parts.append(inline("".join(lines[start:end])))
        cursor = end
    parts.append("".join(lines[cursor:]))
    result = "".join(parts)
    if not registry.used:
        return result
    heading = "引用来源" if re.search(r"[\u4e00-\u9fff]", content) else "Sources"
    references = []
    for identifier, record in registry.records.items():
        if identifier not in registry.used:
            continue
        label = escape(record.get("label") or record.get("title") or identifier)
        locator = str(record.get("locator") or record.get("url") or "")
        href = _safe_link(locator)
        location = (
            f"<{quote(href, safe='/:?&=#@!$%()*+,;[]')}>" if href else escape(locator)
        )
        references.append(
            f"- [{registry.numbers[identifier]}] {label}"
            + (f" — {location}" if location else "")
        )
    return result.rstrip() + f"\n\n## {heading}\n\n" + "\n".join(references) + "\n"


def _clean(value: object) -> str:
    return _INVALID_XML.sub("", str(value or ""))


def _safe_link(value: str) -> str | None:
    value = value.strip()
    if any(ord(character) < 32 for character in value):
        return None
    try:
        url = urlsplit(value)
        if (
            url.scheme in {"http", "https"}
            and url.hostname
            and not url.username
            and not url.password
        ):
            return value
        if url.scheme == "mailto" and url.path and "@" in url.path:
            return value
    except ValueError:
        pass
    return None


def _set_font(properties, *, monospace: bool = False) -> None:
    fonts = properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.insert(0, fonts)
    # Theme font attributes take precedence over explicit faces in Word/LibreOffice.
    for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn(f"w:{attribute}"), None)
    fonts.set(qn("w:ascii"), "Consolas" if monospace else "Arial")
    fonts.set(qn("w:hAnsi"), "Consolas" if monospace else "Arial")
    fonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")


def _configure(document: DocumentObject) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.85)
    section.left_margin = section.right_margin = Inches(1)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string("222222")
    normal.paragraph_format.line_spacing = 1.2
    normal.paragraph_format.space_after = Pt(8)
    _set_font(normal.element.get_or_add_rPr())
    for name, size in (
        ("Title", 26),
        ("Heading 1", 19),
        ("Heading 2", 15),
        ("Heading 3", 12),
    ):
        style = document.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.underline = False
        style.font.color.rgb = RGBColor(0, 0, 0)
        paragraph_properties = style.element.get_or_add_pPr()
        for border in list(paragraph_properties.findall(qn("w:pBdr"))):
            paragraph_properties.remove(border)
        style.paragraph_format.space_before = Pt(16 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True
        _set_font(style.element.get_or_add_rPr())
    for level in range(4, 10):
        style = document.styles[f"Heading {level}"]
        style.font.color.rgb = RGBColor(0, 0, 0)
        _set_font(style.element.get_or_add_rPr())


def _add_text(
    paragraph: Paragraph,
    text: str,
    *,
    bold=False,
    italic=False,
    strike=False,
    code=False,
) -> None:
    run = paragraph.add_run(_clean(text))
    run.bold = bold
    run.italic = italic
    run.font.strike = strike
    _set_font(run._r.get_or_add_rPr(), monospace=code)
    if code:
        run.font.size = Pt(9)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "F1F2F4")
        run._r.get_or_add_rPr().append(shading)


def _hyperlink(
    paragraph: Paragraph, label: str, href: str, *, bold=False, italic=False
) -> None:
    safe = _safe_link(href)
    if safe is None:
        _add_text(paragraph, label, bold=bold, italic=italic)
        return
    from docx.opc.constants import RELATIONSHIP_TYPE

    relation = paragraph.part.relate_to(
        safe, RELATIONSHIP_TYPE.HYPERLINK, is_external=True
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    _set_font(properties)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2451D1")
    properties.append(color)
    if bold:
        properties.append(OxmlElement("w:b"))
    if italic:
        properties.append(OxmlElement("w:i"))
    run.append(properties)
    text = OxmlElement("w:t")
    text.set(qn("xml:space"), "preserve")
    text.text = _clean(label)
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


class _Citations:
    def __init__(self, records: Sequence[Mapping[str, Any]]) -> None:
        self.records: dict[str, Mapping[str, Any]] = {}
        self.used: set[str] = set()
        self.ambiguous: set[str] = set()
        self.numbers: dict[str, int] = {}
        for record in records:
            identifier = record.get("citation_id")
            if not isinstance(identifier, str) or not _CITATION_ID.fullmatch(
                identifier
            ):
                continue
            self.numbers.setdefault(identifier, len(self.numbers) + 1)
            if identifier in self.ambiguous:
                continue
            previous = self.records.get(identifier)
            if previous is not None and previous != record:
                # Conflicting duplicate IDs are not evidence. Identical event
                # replay is harmless, but choosing the first conflicting source
                # would confidently attribute a claim to an arbitrary document.
                self.ambiguous.add(identifier)
                self.records.pop(identifier)
            else:
                self.records[identifier] = record

    def unresolved(self, identifier: str) -> str:
        reason = "来源不明确" if identifier in self.ambiguous else "来源未提供"
        return f"[{reason}: {identifier}]"

    def resolve(self, text: str) -> str:
        def citation(match: re.Match[str]) -> str:
            identifier = match.group(1)
            if identifier not in self.records:
                return self.unresolved(identifier)
            self.used.add(identifier)
            return f"[{self.numbers[identifier]}]"

        return _CITATION.sub(citation, text)

    def append_sources(self, document: DocumentObject) -> None:
        if not self.used:
            return
        document.add_heading(
            "引用来源"
            if any(
                re.search(r"[\u4e00-\u9fff]", str(record.get("label", "")))
                for record in self.records.values()
            )
            else "Sources",
            level=1,
        )
        for identifier in self.records:
            if identifier not in self.used:
                continue
            number = self.numbers[identifier]
            record = self.records[identifier]
            label = str(
                record.get("label")
                or record.get("title")
                or record.get("source")
                or identifier
            )
            paragraph = document.add_paragraph()
            _add_text(paragraph, f"[{number}] {label}")
            page = record.get("page")
            if page is not None:
                _add_text(paragraph, f" · p. {page}")
            locator = str(record.get("locator") or record.get("url") or "")
            if locator:
                paragraph.add_run(" — ")
                _hyperlink(paragraph, locator, locator)


def _inline(
    paragraph: Paragraph, tokens: Sequence[Token], citations: _Citations
) -> None:
    bold = italic = strike = False
    href: str | None = None
    for token in tokens:
        if token.type in {"strong_open", "strong_close"}:
            bold = token.type.endswith("open")
        elif token.type in {"em_open", "em_close"}:
            italic = token.type.endswith("open")
        elif token.type in {"s_open", "s_close"}:
            strike = token.type.endswith("open")
        elif token.type == "link_open":
            href = token.attrGet("href") or ""
        elif token.type == "link_close":
            href = None
        elif token.type in {"softbreak", "hardbreak"}:
            paragraph.add_run().add_break()
        elif token.type == "image":
            # Preserve the description as editable text; never fetch remote media.
            _add_text(paragraph, f"[Image: {token.content}]", italic=True)
        elif token.type in {"text", "code_inline", "html_inline"}:
            text = (
                token.content
                if token.type == "code_inline" or href is not None
                else citations.resolve(token.content)
            )
            if href is not None:
                _hyperlink(paragraph, text, href, bold=bold, italic=italic)
            else:
                _add_text(
                    paragraph,
                    text,
                    bold=bold,
                    italic=italic,
                    strike=strike,
                    code=token.type == "code_inline",
                )


def _numbering(
    document: DocumentObject, *, ordered: bool, depth: int, start: int
) -> int:
    root = document.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in root.findall(qn("w:abstractNum"))
    ]
    abstract_id = max(abstract_ids, default=-1) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for tag, value in (
        ("start", str(start)),
        ("numFmt", "decimal" if ordered else "bullet"),
        ("lvlText", "%1." if ordered else "•"),
    ):
        item = OxmlElement(f"w:{tag}")
        item.set(qn("w:val"), value)
        level.append(item)
    properties = OxmlElement("w:pPr")
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), str(360 * depth))
    indent.set(qn("w:hanging"), "240")
    properties.append(indent)
    level.append(properties)
    abstract.append(level)
    root.append(abstract)
    return int(root.add_num(abstract_id).numId)


def _apply_number(paragraph: Paragraph, number_id: int) -> None:
    properties = paragraph._p.get_or_add_pPr()
    number = OxmlElement("w:numPr")
    for tag, value in (("ilvl", "0"), ("numId", str(number_id))):
        item = OxmlElement(f"w:{tag}")
        item.set(qn("w:val"), value)
        number.append(item)
    properties.append(number)
    paragraph.paragraph_format.space_after = Pt(4)


def _table(
    document: DocumentObject, tokens: Sequence[Token], citations: _Citations
) -> None:
    rows: list[list[list[Token]]] = []
    row: list[list[Token]] = []
    for token in tokens:
        if token.type == "tr_open":
            row = []
        elif token.type in {"th_open", "td_open"}:
            row.append([])
        elif token.type == "inline" and row:
            row[-1].extend(token.children or [])
        elif token.type == "tr_close":
            rows.append(row)
    if not rows:
        return
    columns = max(len(row) for row in rows)
    if columns > 40 or len(rows) * columns > 20_000:
        raise DocumentExportError("Artifact table exceeds the Word export size limit")
    table = document.add_table(rows=len(rows), cols=columns)
    table.autofit = False
    weights = []
    for column_index in range(columns):
        lengths = [
            sum(len(token.content) for token in row[column_index])
            for row in rows
            if column_index < len(row)
        ]
        weights.append(max(10, min(60, sum(lengths) / max(1, len(lengths)))))
    for column, weight in zip(table.columns, weights, strict=True):
        column.width = Inches(6.5 * weight / sum(weights))
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement(f"w:{edge}")
        for name, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            item.set(qn(f"w:{name}"), value)
        borders.append(item)
    properties.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for edge in ("top", "left", "bottom", "right"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:w"), "100" if edge in {"top", "bottom"} else "120")
        item.set(qn("w:type"), "dxa")
        margins.append(item)
    properties.append(margins)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for row_index, row in enumerate(rows):
        for column_index, cell_tokens in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.width = table.columns[column_index].width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            _inline(paragraph, cell_tokens, citations)
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "F3F4F6")
                cell._tc.get_or_add_tcPr().append(shading)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(3)


def _markdown(
    document: DocumentObject, content: str, citations: _Citations, *, title: str
) -> None:
    parser = (
        MarkdownIt("commonmark", {"html": False})
        .enable("table")
        .enable("strikethrough")
    )
    tokens = parser.parse(content)
    if (
        len(tokens) >= 3
        and tokens[0].type == "heading_open"
        and tokens[0].tag == "h1"
        and tokens[1].type == "inline"
        and tokens[1].content.strip() == title.strip()
        and tokens[2].type == "heading_close"
    ):
        tokens = tokens[3:]
    if len(tokens) > 50_000:
        raise DocumentExportError("Artifact has too many document blocks to export")
    paragraph: Paragraph | None = None
    lists: list[dict[str, Any]] = []
    quote_depth = 0
    index = 0
    while index < len(tokens):
        token = tokens[index]
        if token.type in {"bullet_list_open", "ordered_list_open"}:
            lists.append(
                {
                    "number_id": _numbering(
                        document,
                        ordered=token.type == "ordered_list_open",
                        depth=len(lists) + 1,
                        start=int(token.attrGet("start") or 1),
                    ),
                    "pending": False,
                }
            )
        elif token.type in {"bullet_list_close", "ordered_list_close"}:
            lists.pop()
        elif token.type == "list_item_open":
            lists[-1]["pending"] = True
        elif token.type == "blockquote_open":
            quote_depth += 1
        elif token.type == "blockquote_close":
            quote_depth -= 1
        elif token.type in {"paragraph_open", "heading_open"}:
            paragraph = (
                document.add_heading(level=min(int(token.tag[1:]), 6))
                if token.type == "heading_open"
                else document.add_paragraph()
            )
            if lists:
                if lists[-1]["pending"]:
                    _apply_number(paragraph, lists[-1]["number_id"])
                    lists[-1]["pending"] = False
                else:
                    paragraph.paragraph_format.left_indent = Inches(0.25 * len(lists))
            if quote_depth:
                paragraph.paragraph_format.left_indent = Inches(0.3 * quote_depth)
        elif token.type == "inline" and paragraph is not None:
            _inline(paragraph, token.children or [], citations)
        elif token.type in {"paragraph_close", "heading_close"}:
            paragraph = None
        elif token.type in {"fence", "code_block"}:
            code = document.add_paragraph()
            _add_text(code, token.content.rstrip("\n"), code=True)
            code.paragraph_format.line_spacing = 1.05
            code.paragraph_format.space_before = Pt(6)
            code.paragraph_format.space_after = Pt(10)
        elif token.type == "table_open":
            end = index + 1
            while end < len(tokens) and tokens[end].type != "table_close":
                end += 1
            _table(document, tokens[index + 1 : end], citations)
            index = end
        elif token.type == "hr":
            document.add_paragraph("―" * 24)
        index += 1


def export_artifact_docx(
    *,
    title: str,
    content: str,
    media_type: str = "text/markdown",
    citations: Sequence[Mapping[str, Any]] = (),
) -> bytes:
    """Export editable text, headings, lists, tables and links using current-Run citations.

    ``citations`` must already be scoped to the Artifact's Run by the caller. The
    exporter resolves exact citation IDs only and never guesses references by position.
    HTML has a separate download path and is deliberately rejected here.
    """
    if media_type not in {"text/markdown", "text/plain", "application/json"}:
        raise DocumentExportError("This artifact type does not support Word export")
    if len(content) > MAX_EXPORT_CHARS or len(title) > 500:
        raise DocumentExportError("Artifact exceeds the Word export size limit")
    document = Document()
    _configure(document)
    document.core_properties.title = _clean(title)[:255]
    document.core_properties.author = "Alcuin"
    document.add_heading(_clean(title) or "Untitled", level=0)
    source_map = _Citations(citations)
    content = _clean(content)
    if media_type == "text/markdown":
        _markdown(document, content, source_map, title=title)
    elif media_type == "application/json":
        try:
            content = json.dumps(json.loads(content), ensure_ascii=False, indent=2)
        except (ValueError, RecursionError) as exc:
            raise DocumentExportError("Artifact content is not valid JSON") from exc
        _add_text(document.add_paragraph(), content, code=True)
    else:
        for line in content.split("\n"):
            _add_text(document.add_paragraph(), source_map.resolve(line))
    source_map.append_sources(document)
    target = BytesIO()
    document.save(target)
    return target.getvalue()
