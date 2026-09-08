from __future__ import annotations

from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree
import pytest

from alcuin_documents import DocumentExportError, export_artifact_docx
from alcuin_documents.export import MAX_EXPORT_CHARS, export_artifact_markdown


def test_markdown_export_keeps_format_and_portable_sources_without_touching_code():
    content = (
        "## 中文结论\n\n真实 [[cite:s2]] 与 [证据](alcuin-citation:s1)。"
        "未知 [[cite:s9]]，`[[cite:s1]]` 是示例。\n\n"
        "~~~text\n[[cite:s1]]\n~~~\n\n    [[cite:s1]]\n"
    )
    output = export_artifact_markdown(
        content,
        citations=[
            {
                "citation_id": "s1",
                "label": "来源一",
                "locator": "https://example.com/a",
            },
            {
                "citation_id": "s2",
                "label": "知识",
                "locator": "knowledge://document/chunk",
            },
        ],
    )
    assert "真实 [2] 与 [1](<https://example.com/a>)。" in output
    assert "[来源未提供: s9]" in output
    assert "`[[cite:s1]]`" in output
    assert "~~~text\n[[cite:s1]]\n~~~\n\n    [[cite:s1]]" in output
    assert "- [1] 来源一 — <https://example.com/a>" in output
    assert "- [2] 知识 — knowledge://document/chunk" in output


def test_markdown_export_escapes_labels_and_refuses_unsafe_links():
    output = export_artifact_markdown(
        "[[cite:s1]]",
        citations=[
            {
                "citation_id": "s1",
                "label": "[bad](javascript:alert(1))",
                "locator": "javascript:alert(1)",
            },
        ],
    )
    assert "[1](" not in output
    assert r"\[bad\](javascript:alert(1))" in output
    literal = "``a ` [[cite:s1]]`` and `[[cite:s2]]`"
    assert export_artifact_markdown(literal) == literal


@pytest.mark.parametrize(
    "literal",
    [
        "`[[cite:s1]]` and ``code ` [[cite:s1]]``",
        "[Existing [[cite:s1]]](https://example.com/target)",
        "[Existing](https://example.com/?q=[[cite:s1]])",
        "![Example [[cite:s1]]](https://example.com/image.png)",
        '<span title="[[cite:s1]]">example</span>',
        r"\[[cite:s1]]",
        "~~~text\n[[cite:s1]]\n~~~",
        "    [[cite:s1]]",
    ],
)
def test_markdown_citations_preserve_literal_syntax_without_registering_sources(
    literal,
):
    record = {"citation_id": "s1", "label": "Source", "locator": "https://source.test"}
    assert export_artifact_markdown(literal, citations=[record]) == literal


def test_markdown_citations_respect_paragraph_boundaries_and_escaped_backticks():
    record = {"citation_id": "s1", "label": "Source", "locator": "https://source.test"}
    output = export_artifact_markdown(
        "` unmatched\n\nReal [[cite:s1]]\n\nend `\n\n"
        "\\`Real [[cite:s1]]` also [[cite:s1]]\n\n"
        "[Reference][ref]\n\n[ref]: https://example.com/?q=[[cite:s1]]\n",
        citations=[record],
    )
    assert "Real [1](<https://source.test>)" in output
    assert (
        "\\`Real [1](<https://source.test>)` also [1](<https://source.test>)" in output
    )
    assert "[ref]: https://example.com/?q=[[cite:s1]]" in output
    assert output.count("- [1] Source") == 1


def test_conflicting_ids_never_resolve_to_the_first_source_in_markdown_or_docx():
    first = {"citation_id": "s1", "label": "First", "locator": "https://first.test"}
    second = {"citation_id": "s1", "label": "Second", "locator": "https://second.test"}
    third = {"citation_id": "s2", "label": "Valid", "locator": "https://valid.test"}
    records = [first, second, first, third, third]
    content = "Conflicting [[cite:s1]], valid [[cite:s2]], unknown [[cite:s9]]."
    markdown = export_artifact_markdown(content, citations=records)
    _, _, xml = _export(content, citations=records)
    for output in [markdown, _text(xml)]:
        assert "[来源不明确: s1]" in output
        assert "[来源未提供: s9]" in output
        assert "[2] Valid" in output
        assert "First" not in output and "Second" not in output
        assert "first.test" not in output and "second.test" not in output
    assert markdown.count("- [2] Valid") == 1


def test_invalid_ids_are_not_coerced_or_matched_by_position():
    content = "[[cite:1]] [[cite:s1]] [[cite:S1]] [[cite:s1.extra]]"
    records = [
        {"citation_id": 1, "label": "Number"},
        {"id": "s1", "label": "Event alias"},
        {"citation_id": " s1 ", "label": "Whitespace alias"},
    ]
    markdown = export_artifact_markdown(content, citations=records)
    assert "来源未提供: 1" in markdown
    assert "来源未提供: s1" in markdown
    assert "来源未提供: S1" in markdown
    assert "来源未提供: s1.extra" in markdown
    assert "## Sources" not in markdown


def test_markdown_source_labels_cannot_inject_additional_blocks():
    output = export_artifact_markdown(
        "[[cite:s1]]",
        citations=[{"citation_id": "s1", "label": "Title\n\n# Injected\n- list"}],
    )
    assert "\n# Injected" not in output
    assert "- [1] Title # Injected - list" in output


def _export(content: str, **kwargs):
    data = export_artifact_docx(title="Alcuin 工作简报", content=content, **kwargs)
    archive = ZipFile(BytesIO(data))
    document = etree.fromstring(archive.read("word/document.xml"))
    return data, archive, document


def _text(xml) -> str:
    return "".join(xml.itertext())


def test_markdown_exports_editable_headings_numbered_lists_table_and_code() -> None:
    data, archive, xml = _export(
        "# 调研结论\n\n一段 **粗体**、*强调* 与 `code`。\n\n"
        "3. 收集材料\n4. 形成结论\n   - 引用来源\n\n"
        "| 项目 | 结果 |\n| --- | --- |\n| 文档 | 可编辑 |\n\n"
        "```python\nprint('hello')\n```"
    )
    doc = Document(BytesIO(data))
    assert any(
        paragraph.style.name == "Heading 1" and paragraph.text == "调研结论"
        for paragraph in doc.paragraphs
    )
    assert doc.tables[0].cell(1, 1).text == "可编辑"
    assert "print('hello')" in _text(xml)
    assert xml.findall(f".//{qn('w:numPr')}")
    assert xml.findall(f".//{qn('w:tblHeader')}")
    assert xml.findall(f".//{qn('w:tblCellMar')}")
    borders = xml.findall(f".//{qn('w:tblBorders')}/*")
    assert len(borders) == 6
    assert all(border.get(qn("w:color")) == "D9D9D9" for border in borders)
    numbering = etree.fromstring(archive.read("word/numbering.xml"))
    assert any(
        element.get(qn("w:val")) == "3"
        for element in numbering.findall(f".//{qn('w:start')}")
    )
    assert not any(name.startswith("word/media/") for name in archive.namelist())


def test_export_uses_letter_portrait_black_headings_and_cjk_font() -> None:
    _, archive, xml = _export("# 一级\n\n## 二级\n\n中文与English正文。")
    size = xml.find(f".//{qn('w:pgSz')}")
    assert size.get(qn("w:w")) == "12240"
    assert size.get(qn("w:h")) == "15840"
    assert any(
        font.get(qn("w:eastAsia")) == "Noto Sans CJK SC"
        for font in xml.findall(f".//{qn('w:rFonts')}")
    )
    styles = etree.fromstring(archive.read("word/styles.xml"))
    for identifier in ["Title", "Heading1", "Heading2", "Heading3"]:
        style = next(
            node
            for node in styles.findall(qn("w:style"))
            if node.get(qn("w:styleId")) == identifier
        )
        assert style.find(f"{qn('w:rPr')}/{qn('w:color')}").get(qn("w:val")) == "000000"
        assert style.find(f"{qn('w:pPr')}/{qn('w:pBdr')}") is None
        fonts = style.find(f"{qn('w:rPr')}/{qn('w:rFonts')}")
        assert not any("theme" in key.lower() for key in fonts.attrib)


def test_known_citations_keep_registry_numbering_dedupe_sources_and_show_unknowns() -> (
    None
):
    _, _, xml = _export(
        "先用第二来源 [[cite:s2]]，后用第一来源 [[cite:s1]]。重复 [[cite:s2]]；未知 [[cite:s9]]。",
        citations=[
            {
                "citation_id": "s1",
                "label": "来源一",
                "locator": "https://example.com/one",
            },
            {
                "citation_id": "s2",
                "label": "来源二",
                "locator": "https://example.com/two",
            },
            {
                "citation_id": "s3",
                "label": "未引用来源",
                "locator": "https://example.com/unused",
            },
        ],
    )
    text = _text(xml)
    assert (
        "先用第二来源 [2]，后用第一来源 [1]。重复 [2]；未知 [来源未提供: s9]。" in text
    )
    assert text.count("[1] 来源一") == 1
    assert text.count("[2] 来源二") == 1
    assert "未引用来源" not in text
    assert "[[cite:" not in text
    assert text.index("[1] 来源一") < text.index("[2] 来源二")


def test_citation_ids_never_fall_back_to_event_ids_or_guess_pdf_pages() -> None:
    _, _, xml = _export(
        "证据 [[cite:s1]] 与 [[cite:evt1]]。",
        citations=[
            {
                "citation_id": "s1",
                "label": "Manual.pdf",
                "metadata": {"chunk_index": 12, "document_id": "doc_1"},
            },
            {"id": "evt1", "label": "Unidentified event"},
        ],
    )
    text = _text(xml)
    assert "[来源未提供: evt1]" in text
    assert "p. 12" not in text
    assert "Manual.pdf" in text
    assert "Unidentified event" not in text


def test_external_images_are_not_loaded_and_only_safe_hyperlinks_are_created() -> None:
    _, archive, xml = _export(
        "[Public source](https://example.com/source)\n\n"
        "![Remote chart](https://example.com/image.png)\n\n"
        "[unsafe](file:///etc/passwd)\n\n<script>alert('x')</script>",
    )
    relationships = etree.fromstring(archive.read("word/_rels/document.xml.rels"))
    targets = [
        node.get("Target")
        for node in relationships
        if node.get("TargetMode") == "External"
    ]
    assert targets == ["https://example.com/source"]
    assert "[Image: Remote chart]" in _text(xml)
    assert "<script>alert('x')</script>" in _text(xml)
    assert not any(name.startswith("word/media/") for name in archive.namelist())


def test_docx_existing_link_label_is_not_rewritten_into_a_misdirected_citation():
    _, _, xml = _export(
        "[Existing [[cite:s1]]](https://other.test)",
        citations=[
            {"citation_id": "s1", "label": "Evidence", "locator": "https://source.test"}
        ],
    )
    assert "Existing [[cite:s1]]" in _text(xml)
    assert "Evidence" not in _text(xml)


def test_plain_text_does_not_turn_markdown_characters_into_formatting() -> None:
    data, _, xml = _export(
        "# Plain title\n**literal text** [[cite:s1]]", media_type="text/plain"
    )
    document = Document(BytesIO(data))
    assert document.paragraphs[1].text == "# Plain title"
    assert "**literal text** [来源未提供: s1]" in _text(xml)
    assert not any(
        paragraph.style.name == "Heading 1" for paragraph in document.paragraphs
    )


def test_json_exports_pretty_editable_content_and_rejects_invalid_json() -> None:
    _, _, xml = _export('{"name":"中文","items":[1,2]}', media_type="application/json")
    assert '"name": "中文"' in _text(xml)
    assert '"items": [' in _text(xml)
    with pytest.raises(DocumentExportError, match="not valid JSON"):
        _export("{bad json}", media_type="application/json")


def test_export_strips_xml_control_characters_and_enforces_supported_types_and_size() -> (
    None
):
    _, _, xml = _export("Safe\x00 content\x0b.\ud800")
    assert "Safe content." in _text(xml)
    with pytest.raises(DocumentExportError, match="does not support Word"):
        _export("<h1>HTML has a separate download path</h1>", media_type="text/html")
    with pytest.raises(DocumentExportError, match="size limit"):
        _export("x" * (MAX_EXPORT_CHARS + 1))


def test_matching_artifact_title_is_not_duplicated_in_document_body() -> None:
    _, _, xml = _export("# Alcuin 工作简报\n\nThe actual content.")
    assert _text(xml).count("Alcuin 工作简报") == 1
